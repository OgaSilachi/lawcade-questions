#!/usr/bin/env python3
"""Measure the visual and verbal DNA of approved firm documents.

Usage:
    python3 fingerprint.py approved1.docx approved2.docx ... [--json out.json]

Prints a Markdown report: what the documents actually do (fonts, sizes,
spacing, alignment, numbering, tables, sentence length, defined terms),
document by document and as a consensus. DISTILL mode reads this report
before it proposes any change to firm.json.
"""
import argparse
import json
import re
import statistics as stats
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z\"“(])")
DEFINED = re.compile(r"\((?:the |each, a |collectively, the )?[\"“][A-Z][^\"”]{1,40}[\"”]\)")
TYPED_NUM = re.compile(r"^\s*(?:[IVXLC]+\.|[A-Z]\.|\d+(?:\.\d+)*\.?|\([a-z0-9]+\))\s+\S")


def style_chain(style, getter):
    while style is not None:
        if getattr(style, "element", None) is None:
            return None
        v = getter(style)
        if v is not None:
            return v
        style = style.base_style
    return None


def run_font(run, para):
    rf = run._r.rPr.find(qn("w:rFonts")) if run._r.rPr is not None else None
    if rf is not None and rf.get(qn("w:ascii")):
        return rf.get(qn("w:ascii"))
    return style_chain(para.style, lambda s: s.font.name) or "(theme default)"


def run_size(run, para):
    if run.font.size:
        return run.font.size.pt
    return style_chain(para.style, lambda s: s.font.size.pt if s.font.size else None)


def para_value(para, attr):
    v = getattr(para.paragraph_format, attr)
    if v is None:
        v = style_chain(para.style, lambda s: getattr(s.paragraph_format, attr))
    return v


def numbering_formats(doc):
    """Formats of the lists the document actually uses, not every list in the template."""
    out = Counter()
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return out
    used = {el.get(qn("w:val")) for el in doc.element.body.iter(qn("w:numId"))}
    for st in doc.styles:
        if st.element.pPr is not None:
            used |= {el.get(qn("w:val")) for el in st.element.pPr.iter(qn("w:numId"))}
    abstract_ids = {n.find(qn("w:abstractNumId")).get(qn("w:val"))
                    for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) in used}
    lvls = [lvl for a in numbering.findall(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) in abstract_ids for lvl in a.findall(qn("w:lvl"))]
    for lvl in lvls:
        fmt = lvl.find(qn("w:numFmt"))
        txt = lvl.find(qn("w:lvlText"))
        if fmt is not None and txt is not None and int(lvl.get(qn("w:ilvl"))) < 3:
            out[f"L{lvl.get(qn('w:ilvl'))}:{fmt.get(qn('w:val'))} '{txt.get(qn('w:val'))}'"] += 1
    return out


def heading_candidates(doc):
    """Short paragraphs that look like headings, with how they are dressed."""
    looks = Counter()
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t or len(t) > 90 or t.endswith((",", ";")):
            continue
        bold = any(r.bold for r in p.runs if r.text.strip()) or bool(style_chain(p.style, lambda s: s.font.bold))
        under = any(r.underline for r in p.runs if r.text.strip())
        caps = t.isupper()
        small = any(r.font.small_caps for r in p.runs) or bool(style_chain(p.style, lambda s: s.font.small_caps))
        if not (bold or under or caps or small or (p.style is not None and p.style.name.lower().startswith("heading"))):
            continue
        align = para_value(p, "alignment")
        tags = [str(align).split(".")[-1].split(" ")[0] if align is not None else "LEFT"]
        tags += [n for n, f in (("bold", bold), ("underline", under), ("ALLCAPS", caps), ("smallcaps", small)) if f]
        looks[" + ".join(tags)] += 1
    return looks


def fingerprint(path):
    doc = Document(path)
    sec = doc.sections[0]
    fonts, sizes, aligns, leading, after, first = Counter(), Counter(), Counter(), Counter(), Counter(), Counter()
    colors = Counter()
    sentences, typed, emdash, defined = [], 0, 0, 0
    body_chars = 0
    for p in doc.paragraphs:
        t = p.text
        if not t.strip():
            continue
        for r in p.runs:
            n = len(r.text)
            fonts[run_font(r, p)] += n
            s = run_size(r, p)
            if s:
                sizes[s] += n
            if r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) not in ("000000",):
                colors[str(r.font.color.rgb)] += n
        emdash += t.count("—") + t.count("–")
        defined += len(DEFINED.findall(t))
        if len(t) > 120:  # body paragraph
            body_chars += len(t)
            a = para_value(p, "alignment")
            aligns[str(a).split(".")[-1].split(" ")[0] if a is not None else "LEFT"] += 1
            ls = para_value(p, "line_spacing")
            rule = para_value(p, "line_spacing_rule")
            if ls is not None:
                leading[f"{round(ls.pt, 1)}pt exact" if hasattr(ls, "pt") else f"{ls}x"] += 1
            sa = para_value(p, "space_after")
            after[f"{round(sa.pt)}pt" if sa is not None else "0pt"] += 1
            fi = para_value(p, "first_line_indent")
            first[f"{round(fi.inches, 2)}in" if fi is not None else "0in"] += 1
            for s in SENT.split(t):
                w = len(s.split())
                if w > 2:
                    sentences.append(w)
        if TYPED_NUM.match(t) and p._p.pPr is not None and p._p.pPr.numPr is None \
                and style_chain(p.style, lambda s: s.element.pPr.numPr if s.element.pPr is not None else None) is None:
            typed += 1
    tables = doc.tables
    bordered = 0
    for tb in tables:
        b = tb._tbl.tblPr.find(qn("w:tblBorders"))
        if b is not None and any(el.get(qn("w:val")) not in ("nil", "none") for el in b):
            bordered += 1
    text = "\n".join(p.text for p in doc.paragraphs)
    return {
        "file": Path(path).name,
        "margins_in": [round(x.inches, 2) for x in (sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin)],
        "fonts_by_char": top(fonts),
        "sizes_by_char": top(sizes),
        "colors_by_char": top(colors),
        "body_alignment": top(aligns),
        "body_leading": top(leading),
        "body_space_after": top(after),
        "body_first_line_indent": top(first),
        "heading_dress": top(heading_candidates(doc), 6),
        "numbering_levels": top(numbering_formats(doc), 6),
        "typed_numbering_paragraphs": typed,
        "tables": len(tables),
        "tables_with_borders": bordered,
        "avg_sentence_words": round(stats.mean(sentences), 1) if sentences else None,
        "p90_sentence_words": sorted(sentences)[int(len(sentences) * 0.9)] if sentences else None,
        "em_or_en_dashes": emdash,
        "defined_terms": defined,
        "has_signature_page": bool(re.search(r"IN WITNESS|Signature Page|By:\s*_{3,}", text, re.I)),
        "has_summary_table_page1": bool(tables) and tables[0]._tbl in [e for e in doc.element.body.iterchildren()][:25],
    }


def top(counter, n=4):
    total = sum(counter.values()) or 1
    return [[str(k), round(100 * v / total)] for k, v in counter.most_common(n)]


def consensus(fps, key):
    agg = Counter()
    for fp in fps:
        for k, pct in fp[key]:
            agg[k] += pct
    return top(agg, 3)


def report(fps):
    out = ["# Style Fingerprint", "", f"Documents measured: {len(fps)}", ""]
    out += ["## Consensus (what the approved set agrees on)", "", "| Trait | Leading values (share) |", "|---|---|"]
    for key in ("fonts_by_char", "sizes_by_char", "body_alignment", "body_leading", "body_space_after",
                "body_first_line_indent", "heading_dress", "numbering_levels", "colors_by_char"):
        vals = ", ".join(f"{k} ({v}%)" for k, v in consensus(fps, key)) or "none"
        out.append(f"| {key.replace('_', ' ')} | {vals} |")
    nums = [f["avg_sentence_words"] for f in fps if f["avg_sentence_words"]]
    out.append(f"| avg sentence words | {round(stats.mean(nums), 1) if nums else 'n/a'} |")
    out.append(f"| em/en dashes (total) | {sum(f['em_or_en_dashes'] for f in fps)} |")
    out.append(f"| typed-numbering paragraphs (total) | {sum(f['typed_numbering_paragraphs'] for f in fps)} |")
    out.append(f"| docs with signature page | {sum(f['has_signature_page'] for f in fps)} of {len(fps)} |")
    out += ["", "## Per document", ""]
    for f in fps:
        out.append(f"### {f['file']}")
        for k, v in f.items():
            if k != "file":
                out.append(f"- **{k}**: {v}")
        out.append("")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docs", nargs="+")
    ap.add_argument("--json")
    a = ap.parse_args()
    fps = []
    for d in a.docs:
        try:
            fps.append(fingerprint(d))
        except Exception as e:  # keep going; one bad file should not kill the run
            print(f"Skipped {d}: {e}", file=sys.stderr)
    if not fps:
        sys.exit("No readable .docx files.")
    if a.json:
        Path(a.json).write_text(json.dumps(fps, indent=2))
    print(report(fps))


if __name__ == "__main__":
    main()
