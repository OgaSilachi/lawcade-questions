#!/usr/bin/env python3
"""Audit a .docx against the Silachi Prestige Standard.

Usage:
    python3 audit.py draft.docx [--firm ../assets/firm.json] [--voice-only]

Exit code 0 when every hard check passes, 1 otherwise. Use --voice-only for
court filings and letters, where the Red Book governs layout and this
standard governs only the voice rules (Section 9).
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z\"“(])")
TYPED_NUM = re.compile(r"^\s*(?:[IVXLC]+\.|[A-Z]\.|\d+\.)\s+\S")
PLACEHOLDER = re.compile(r"\[\[|\{\{|\[VERIFY|\[[A-Z][A-Z .,/]{2,}\]")


def all_text_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def has_numbering(p):
    s = p.style
    if p._p.pPr is not None and p._p.pPr.numPr is not None:
        return True
    while s is not None:
        if s.element.pPr is not None and s.element.pPr.numPr is not None:
            return True
        s = s.base_style
    return False


def fonts_in(doc):
    found = set()
    for el in doc.element.iter(qn("w:rFonts")):
        v = el.get(qn("w:ascii"))
        if v:
            found.add(v)
    for st in doc.styles:
        el = st.element.find(qn("w:rPr"))
        if el is not None:
            rf = el.find(qn("w:rFonts"))
            if rf is not None and rf.get(qn("w:ascii")) and st.name.startswith("SH"):
                found.add(rf.get(qn("w:ascii")))
    return found


def colors_in(doc):
    return {el.get(qn("w:val")).upper() for el in doc.element.iter(qn("w:color"))
            if el.get(qn("w:val")) not in (None, "auto")}


def audit(path, firm, voice_only):
    doc = Document(path)
    results = []  # (status, check, detail)

    def check(ok, name, detail="", soft=False):
        results.append(("PASS" if ok else ("WARN" if soft else "FAIL"), name, detail))

    text_pars = [p for p in all_text_paragraphs(doc) if p.text.strip()]
    full = "\n".join(p.text for p in text_pars)

    # ---- voice (always) ----
    dashes = [c for c in firm["voice"]["banned_chars"] if c in full]
    check(not dashes, "No em or en dashes", f"found {sum(full.count(c) for c in dashes)}" if dashes else "")
    banned = [b for b in firm["voice"]["banned_phrases"] if re.search(rf"\b{re.escape(b)}\b", full, re.I)]
    check(not banned, "No banned phrases", ", ".join(banned))
    lens = [len(s.split()) for p in text_pars if len(p.text) > 80 for s in SENT.split(p.text) if len(s.split()) > 2]
    if lens:
        avg = sum(lens) / len(lens)
        long_ = [n for n in lens if n > firm["voice"]["max_sentence_words"]]
        check(avg <= firm["voice"]["max_avg_sentence_words"], "Average sentence length",
              f"{avg:.1f} words (limit {firm['voice']['max_avg_sentence_words']})", soft=True)
        check(not long_, "No sentence over the maximum", f"{len(long_)} over {firm['voice']['max_sentence_words']} words", soft=True)
    ph = sorted(set(PLACEHOLDER.findall(full)))
    check(not ph, "No unresolved placeholders", ", ".join(ph[:8]))

    if voice_only:
        return results

    # ---- layout ----
    allowed = {firm["fonts"]["serif"], firm["fonts"]["sans"]}
    extra = fonts_in(doc) - allowed
    check(not extra, "Only the two house typefaces", ", ".join(sorted(extra)))
    palette = {v.upper() for v in firm["colors"].values()} | {"000000"}
    stray = colors_in(doc) - palette
    check(not stray, "Only house colors", ", ".join(sorted(stray)))

    sec = doc.sections[0]
    m = firm["page"]
    ok = (abs(sec.left_margin.inches - m["margin_side_in"]) < 0.02 and abs(sec.right_margin.inches - m["margin_side_in"]) < 0.02
          and abs(sec.top_margin.inches - m["margin_top_in"]) < 0.02 and abs(sec.bottom_margin.inches - m["margin_bottom_in"]) < 0.02)
    check(ok, "Margins", f"{sec.top_margin.inches:.2f}/{sec.right_margin.inches:.2f}/{sec.bottom_margin.inches:.2f}/{sec.left_margin.inches:.2f}")
    check(len(doc.sections) == 1, "One section", f"{len(doc.sections)} sections")

    typed = [p.text[:40] for p in doc.paragraphs if TYPED_NUM.match(p.text) and not has_numbering(p)]
    check(not typed, "No typed numbering", "; ".join(typed[:3]))

    articles = [p for p in doc.paragraphs if p.style.name == "SH Article"]
    check(bool(articles), "Articles use the SH Article style", f"{len(articles)} found")
    check(all(has_numbering(p) for p in articles), "Articles carry the locked list")

    centered_body = [p.text[:40] for p in doc.paragraphs
                     if len(p.text) > 120 and p.alignment is not None and "CENTER" in str(p.alignment)]
    check(not centered_body, "No centered body text", "; ".join(centered_body[:2]))
    underlined = [p.text[:40] for p in doc.paragraphs if p.style.name.startswith("SH") and any(r.underline for r in p.runs)]
    check(not underlined, "No underlined headings", "; ".join(underlined[:2]))

    body = doc.element.body
    empties, run = 0, 0
    for p in doc.paragraphs:
        run = run + 1 if not p.text.strip() and not p._p.xpath(".//w:br") else 0
        empties = max(empties, run)
    check(empties < 2, "No stacked empty paragraphs", f"max run {empties}")

    for t in doc.tables:
        for el in t._tbl.iter(qn("w:insideV")):
            if el.get(qn("w:val")) not in ("nil", "none"):
                check(False, "No vertical table borders")
                break
    sig_tables = [t for t in doc.tables if "By:" in "\n".join(c.text for r in t.rows for c in r.cells)
                  or "____" in "\n".join(c.text for r in t.rows for c in r.cells)]
    check(bool(sig_tables), "Signature page present", soft=True)
    for t in sig_tables:
        check(all(r._tr.trPr is not None and r._tr.trPr.find(qn("w:cantSplit")) is not None for r in t.rows),
              "Signature rows cannot split")
    check("Signature page follows" in full, "[Signature page follows.] line", soft=True)

    footer_xml = "".join(f._element.xml for s in doc.sections for f in (s.footer, s.first_page_footer))
    check("NUMPAGES" in footer_xml and "PAGE" in footer_xml, "Live Page X of Y footer")
    return results


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--firm", default=str(HERE.parent / "assets" / "firm.json"))
    ap.add_argument("--voice-only", action="store_true")
    a = ap.parse_args()
    firm = json.loads(Path(a.firm).read_text())
    res = audit(a.docx, firm, a.voice_only)
    print(f"# House Style Audit: {Path(a.docx).name} ({firm['standard']} v{firm['version']})\n")
    print("| Result | Check | Detail |\n|---|---|---|")
    for status, name, detail in res:
        print(f"| {status} | {name} | {detail} |")
    fails = sum(1 for r in res if r[0] == "FAIL")
    print(f"\n{len(res) - fails} of {len(res)} checks clear. {'CLEAR' if not fails else 'NOT CLEAR'}.")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
